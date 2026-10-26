"""Importación masiva de pedidos (RF-004 / US-004).

Estrategia de persistencia (contrato documental):
- **Atómica (HTTP 400, no se persiste ningún pedido):** archivo vacío, encabezado ausente,
  columnas obligatorias faltantes, más de 150 filas de datos, formato no reconocido o Excel corrupto.
- **Parcial (HTTP 200, informe por fila):** filas válidas se guardan; filas inválidas se rechazan
  con línea, campo y motivo (RF-004 Ruta Infeliz). US-004 Esc. 2 se interpreta como rechazo del
  archivo cuando falla la **estructura** (columnas), no cuando una fila aislada tiene errores.
"""

from __future__ import annotations

import base64
import csv
import io
from datetime import date, datetime, time
from typing import Any
from zoneinfo import ZoneInfo

from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.geo import point_wkt
from app.models.domain import Pedido
from app.models.enums import EstadoPedido, EstadoUbicacion
from app.models.operations import Jornada
from app.services.geocode import (
    MSG_INVALID_COORDS,
    MSG_NOT_FOUND,
    MSG_OUT_OF_AREA,
    MSG_OUT_OF_VALLEY,
    default_altitude,
    parse_coordinates,
    resolve_from_address,
)
from app.services.operating_area import point_in_operating_area
from app.services.order_location import mark_pending_validation
from app.services.tracking_tokens import new_tracking_token

LIMA = ZoneInfo("America/Lima")
MAX_IMPORT_ROWS = 150
_LAT = (-12.30, -11.70)
_LON = (-75.50, -74.90)

_HEADER = {
    "codigo": "codigo",
    "codigo_pedido": "codigo",
    "cliente": "cliente",
    "cliente_nombre": "cliente",
    "direccion": "direccion",
    "direccion_referencia": "direccion",
    "peso_kg": "peso_kg",
    "peso": "peso_kg",
    "hora_inicio": "hora_inicio",
    "ventana_inicio": "hora_inicio",
    "hora_fin": "hora_fin",
    "ventana_fin": "hora_fin",
    "lat": "lat",
    "latitud": "lat",
    "lon": "lon",
    "longitud": "lon",
    "altitud": "altitud",
    "altitud_msnm": "altitud",
    "volumen_m3": "volumen_m3",
    "volumen": "volumen_m3",
}
_REQUIRED = ("codigo", "cliente", "direccion", "peso_kg", "hora_inicio", "hora_fin")


class ImportRow(BaseModel):
    linea: int
    codigo_pedido: str = ""
    cliente_nombre: str = ""
    direccion_referencia: str = ""
    peso_kg: float | None = None
    hora_inicio: str = ""
    hora_fin: str = ""
    lat: float | None = None
    lon: float | None = None
    altitud_msnm: float | None = None
    resultado: str
    motivo: str | None = None
    campo: str | None = None
    pedido_id: int | None = None


class ImportReport(BaseModel):
    fecha: date
    aceptados: int = 0
    rechazados: int = 0
    pendientes_punto: int = 0
    total_filas: int = 0
    resumen: str = ""
    filas: list[ImportRow] = Field(default_factory=list)


def parse_csv(raw: str) -> tuple[list[dict[str, str]], str | None]:
    """Normaliza el CSV. El segundo valor es un error de archivo, no de una fila."""
    text = raw.lstrip("\ufeff").strip()
    if not text:
        return [], "El archivo no tiene filas."
    try:
        reader = csv.DictReader(io.StringIO(text))
    except csv.Error:
        return [], "El archivo CSV está malformado."
    if reader.fieldnames is None:
        return [], "El archivo no tiene encabezado."
    return _rows_from_headers(reader.fieldnames, reader)


def parse_xlsx(raw: bytes) -> tuple[list[dict[str, str]], str | None]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("El servidor no tiene soporte Excel instalado.") from exc
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        return [], "El archivo Excel no se puede leer o está corrupto."
    sheet = workbook.active
    if sheet is None:
        return [], "El archivo Excel no tiene hoja activa."
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], "El archivo no tiene filas."
    headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
    if not any(headers):
        return [], "El archivo no tiene encabezado."
    dict_rows: list[dict[str, str]] = []
    for values in rows_iter:
        if values is None or not any(v is not None and str(v).strip() for v in values):
            continue
        row_map = {
            headers[index]: ("" if values[index] is None else str(values[index]).strip())
            for index in range(min(len(headers), len(values)))
            if headers[index]
        }
        if any(row_map.values()):
            dict_rows.append(row_map)
    if not dict_rows:
        return [], "El archivo no tiene filas de datos."
    return _rows_from_header_maps(headers, dict_rows)


def _rows_from_headers(fieldnames: list[str], reader: csv.DictReader) -> tuple[list[dict[str, str]], str | None]:
    columns = {_HEADER.get(_fold_header(name), ""): name for name in fieldnames}
    missing = [name for name in _REQUIRED if not columns.get(name)]
    if missing:
        return [], f"Faltan columnas obligatorias: {', '.join(missing)}."
    rows: list[dict[str, str]] = []
    for raw_row in reader:
        item = {key: (raw_row.get(source) or "").strip() for key, source in columns.items() if source}
        if not any(item.values()):
            continue
        rows.append(item)
    if not rows:
        return [], "El archivo no tiene filas de datos."
    if len(rows) > MAX_IMPORT_ROWS:
        return [], f"El archivo supera el máximo de {MAX_IMPORT_ROWS} pedidos por importación."
    return rows, None


def _rows_from_header_maps(
    fieldnames: list[str],
    mapped_rows: list[dict[str, str]],
) -> tuple[list[dict[str, str]], str | None]:
    columns = {_HEADER.get(_fold_header(name), ""): name for name in fieldnames}
    missing = [name for name in _REQUIRED if not columns.get(name)]
    if missing:
        return [], f"Faltan columnas obligatorias: {', '.join(missing)}."
    rows: list[dict[str, str]] = []
    for raw_row in mapped_rows:
        item = {key: (raw_row.get(source) or "").strip() for key, source in columns.items() if source}
        if not any(item.values()):
            continue
        rows.append(item)
    if not rows:
        return [], "El archivo no tiene filas de datos."
    if len(rows) > MAX_IMPORT_ROWS:
        return [], f"El archivo supera el máximo de {MAX_IMPORT_ROWS} pedidos por importación."
    return rows, None


async def ensure_jornada(session: AsyncSession, fecha: date) -> Jornada:
    found = await session.scalar(select(Jornada).where(Jornada.fecha == fecha))
    if found is not None:
        return found
    created = Jornada(fecha=fecha)
    session.add(created)
    await session.flush()
    return created


async def import_orders(
    session: AsyncSession,
    fecha: date,
    *,
    csv: str | None = None,
    xlsx_bytes: bytes | None = None,
) -> ImportReport:
    if (csv is None) == (xlsx_bytes is None):
        raise ValueError("Indica un archivo CSV o Excel (.xlsx), no ambos.")
    if xlsx_bytes is not None:
        parsed, error = parse_xlsx(xlsx_bytes)
    else:
        parsed, error = parse_csv(csv or "")
    if error is not None:
        raise ValueError(error)
    jornada = await ensure_jornada(session, fecha)
    existing = set(await session.scalars(select(Pedido.codigo_pedido)))
    seen: set[str] = set()
    report = ImportReport(fecha=fecha, total_filas=len(parsed))
    for index, item in enumerate(parsed, start=2):
        row, order = _classify(fecha, index, item, existing, seen)
        report.filas.append(row)
        if order is None:
            if row.resultado == "rechazado":
                report.rechazados += 1
            continue
        order.jornada_id = jornada.id
        if (
            order.estado_ubicacion == EstadoUbicacion.GEOCODIFICADO
            and row.lat is not None
            and row.lon is not None
        ):
            if not await point_in_operating_area(session, row.lon, row.lat):
                mark_pending_validation(order)
                row.lat = None
                row.lon = None
                row.altitud_msnm = None
                row.resultado = "pendiente_punto"
                row.motivo = MSG_OUT_OF_AREA
                report.pendientes_punto += 1
            else:
                report.aceptados += 1
        else:
            report.pendientes_punto += 1
        session.add(order)
        await session.flush()
        row.pedido_id = order.id
        existing.add(order.codigo_pedido)
        seen.add(order.codigo_pedido)
    report.resumen = _summary(report)
    return report


def decode_xlsx_payload(value: str) -> bytes:
    try:
        raw = base64.b64decode(value, validate=True)
    except Exception as exc:
        raise ValueError("El archivo Excel no está codificado correctamente.") from exc
    if not raw:
        raise ValueError("El archivo Excel está vacío.")
    return raw


def _summary(report: ImportReport) -> str:
    return (
        f"{report.aceptados} aceptados, {report.rechazados} rechazados, "
        f"{report.pendientes_punto} pendientes de punto "
        f"(de {report.total_filas} filas)."
    )


def _reject(row: ImportRow, campo: str, motivo: str) -> tuple[ImportRow, None]:
    row.resultado = "rechazado"
    row.campo = campo
    row.motivo = motivo
    return row, None


def _classify(
    fecha: date,
    linea: int,
    item: dict[str, str],
    existing: set[str],
    seen: set[str],
) -> tuple[ImportRow, Pedido | None]:
    row = ImportRow(
        linea=linea,
        codigo_pedido=item.get("codigo", ""),
        cliente_nombre=item.get("cliente", ""),
        direccion_referencia=item.get("direccion", ""),
        hora_inicio=item.get("hora_inicio", ""),
        hora_fin=item.get("hora_fin", ""),
        resultado="rechazado",
    )
    code = row.codigo_pedido.strip()
    row.codigo_pedido = code
    if not code:
        return _reject(row, "codigo", "Falta el código del pedido.")
    if len(code) < 3:
        return _reject(row, "codigo", "El código es demasiado corto.")
    if code in seen:
        return _reject(row, "codigo", f"El código {code} está duplicado en el archivo.")
    if code in existing:
        return _reject(row, "codigo", f"Ya existe un pedido con el código {code}.")
    if not row.cliente_nombre:
        return _reject(row, "cliente", "Falta el nombre del cliente.")
    if len(row.cliente_nombre) < 2:
        return _reject(row, "cliente", "El nombre del cliente es demasiado corto.")
    if not row.direccion_referencia:
        return _reject(row, "direccion", "Falta la dirección.")
    if len(row.direccion_referencia) < 3:
        return _reject(row, "direccion", "La dirección es demasiado corta.")
    peso_text = item.get("peso_kg", "").strip()
    if not peso_text:
        return _reject(row, "peso_kg", "Falta el peso.")
    try:
        weight = float(peso_text)
    except ValueError:
        return _reject(row, "peso_kg", "El peso no es un número.")
    if weight <= 0:
        return _reject(row, "peso_kg", "El peso tiene que ser mayor que cero.")
    row.peso_kg = weight
    volume: float | None = None
    volume_text = item.get("volumen_m3", "").strip()
    if volume_text:
        try:
            volume = float(volume_text)
        except ValueError:
            return _reject(row, "volumen_m3", "El volumen no es un número.")
        if volume <= 0:
            return _reject(row, "volumen_m3", "El volumen tiene que ser mayor que cero.")
    if not row.hora_inicio.strip():
        return _reject(row, "hora_inicio", "Falta la hora de inicio de la ventana.")
    if not row.hora_fin.strip():
        return _reject(row, "hora_fin", "Falta la hora de fin de la ventana.")
    start = _clock(row.hora_inicio)
    end = _clock(row.hora_fin)
    if start is None:
        return _reject(row, "hora_inicio", "La hora de inicio no es válida (use HH:MM).")
    if end is None:
        return _reject(row, "hora_fin", "La hora de fin no es válida (use HH:MM).")
    if end <= start:
        return _reject(row, "hora_fin", "La ventana horaria debe terminar después del inicio.")
    lat, lon, altitude, located = _point(item, row.direccion_referencia)
    row.lat = lat
    row.lon = lon
    row.altitud_msnm = altitude
    wrote_coordinates = bool(item.get("lat", "").strip() or item.get("lon", "").strip())
    if wrote_coordinates and not located:
        return _reject(row, "lat", MSG_INVALID_COORDS)
    if not located:
        row.resultado = "pendiente_punto"
        row.motivo = MSG_NOT_FOUND
        seen.add(code)
        return row, _pending_order(
            code=code,
            row=row,
            weight=weight,
            volume_m3=volume,
            fecha=fecha,
            start=start,
            end=end,
        )
    if lat is None or lon is None or not (_LAT[0] <= lat <= _LAT[1] and _LON[0] <= lon <= _LON[1]):
        return _reject(row, "lat", MSG_OUT_OF_VALLEY)
    altitude = default_altitude(altitude)
    row.altitud_msnm = altitude
    order = Pedido(
        codigo_pedido=code,
        cliente_nombre=row.cliente_nombre,
        direccion_referencia=row.direccion_referencia,
        ubicacion=point_wkt(lon, lat),
        altitud_msnm=altitude,
        peso_kg=weight,
        volumen_m3=volume,
        ventana_inicio=datetime.combine(fecha, start, tzinfo=LIMA),
        ventana_fin=datetime.combine(fecha, end, tzinfo=LIMA),
        estado=EstadoPedido.PENDING,
        estado_ubicacion=EstadoUbicacion.GEOCODIFICADO,
        token_seguimiento=new_tracking_token(),
    )
    row.resultado = "aceptado"
    row.motivo = "Ubicado por la dirección." if item.get("lat", "") == "" else None
    return row, order


def _pending_order(
    *,
    code: str,
    row: ImportRow,
    weight: float,
    volume_m3: float | None,
    fecha: date,
    start: time,
    end: time,
) -> Pedido:
    return Pedido(
        codigo_pedido=code,
        cliente_nombre=row.cliente_nombre,
        direccion_referencia=row.direccion_referencia,
        ubicacion=None,
        altitud_msnm=None,
        peso_kg=weight,
        volumen_m3=volume_m3,
        ventana_inicio=datetime.combine(fecha, start, tzinfo=LIMA),
        ventana_fin=datetime.combine(fecha, end, tzinfo=LIMA),
        estado=EstadoPedido.PENDING,
        estado_ubicacion=EstadoUbicacion.PENDIENTE_VALIDACION,
        token_seguimiento=new_tracking_token(),
    )


def _point(
    item: dict[str, str],
    address: str,
) -> tuple[float | None, float | None, float | None, bool]:
    lat, lon, altitude, ok = parse_coordinates(
        item.get("lat", ""),
        item.get("lon", ""),
        item.get("altitud", ""),
    )
    if ok and lat is not None and lon is not None:
        return lat, lon, altitude, True
    resolved = resolve_from_address(address)
    if resolved is None:
        return None, None, altitude, False
    return (
        resolved.lat,
        resolved.lon,
        default_altitude(altitude if altitude is not None else resolved.altitude_m),
        True,
    )


def _optional_float(value: str) -> float | None:
    if value == "":
        return None
    try:
        return float(value)
    except ValueError:
        return None


def _clock(value: str) -> time | None:
    try:
        hour, minute = value.strip().split(":", maxsplit=1)
        parsed = time(int(hour), int(minute))
    except (ValueError, AttributeError):
        return None
    return parsed


def _fold_header(name: str | None) -> str:
    return (name or "").strip().casefold()
